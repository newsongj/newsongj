import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.worksheet.page import PageMargins

# 엑셀 파일 읽기
input_file = 'input.xlsx'
wb = openpyxl.load_workbook(input_file)
ws = wb.active

# 새로운 워크북 생성
report_wb = Workbook()
report_ws = report_wb.active

# A4 용지 설정
report_ws.page_setup.paperSize = report_ws.PAPERSIZE_A4
report_ws.page_setup.orientation = report_ws.ORIENTATION_PORTRAIT
report_ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.5, bottom=0.5)

# 폰트 및 스타일 설정
header_font = Font(bold=True, size=14)
content_font = Font(size=12)
title_font = Font(bold=True, size=16)
opinion_font = Font(size=8)
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

# 데이터 읽기
headers = [cell.value for cell in ws[1]]
data = [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]
header_to_index = {str(header).strip(): idx for idx, header in enumerate(headers) if header is not None}

def get_value(row, column_name):
    idx = header_to_index.get(column_name)
    if idx is None:
        return None
    return row[idx]

def get_value_or_index(row, column_name, index):
    value = get_value(row, column_name)
    if value is None and index < len(row):
        return row[index]
    return value

def get_text_or_fallback(row, primary, fallback):
    value = get_value(row, primary)
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return get_value(row, fallback)
    return value

def normalize(value):
    return "" if value is None else value

def apply_row_border(ws, row, start_column, end_column):
    for col in range(start_column, end_column + 1):
        ws.cell(row=row, column=col).border = border

def set_value_cell(ws, row, column, value, font):
    cell = ws.cell(row=row, column=column, value=value)
    cell.font = font
    cell.alignment = Alignment(horizontal='left', vertical='center')
    return cell

# 보고서 작성
for row_idx, row_data in enumerate(data, start=1):
    start_row = (row_idx - 1) * 42 + 1
    
    # 제목
    report_ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=6)
    report_ws.cell(row=start_row, column=1, value="2025 NEWSONG J 그룹배치").font = title_font
    report_ws.cell(row=start_row, column=1).alignment = Alignment(horizontal='center', vertical='center')
    
    report_ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=6)
    report_ws.cell(row=start_row + 1, column=1, value="하나님의 열심이 이루시리라").font = title_font
    report_ws.cell(row=start_row + 1, column=1).alignment = Alignment(horizontal='center', vertical='center')
    
    report_ws.merge_cells(start_row=start_row + 2, start_column=1, end_row=start_row + 2, end_column=6)
    report_ws.cell(row=start_row + 2, column=1, value="그룹원 소견서").font = title_font
    report_ws.cell(row=start_row + 2, column=1).alignment = Alignment(horizontal='center', vertical='center')
    
    # 기본 정보
    report_ws.cell(row=start_row + 4, column=1, value="이름").font = header_font
    set_value_cell(report_ws, start_row + 4, column=2, value=normalize(get_value(row_data, "이름")), font=content_font)
    report_ws.cell(row=start_row + 4, column=3, value="성별").font = header_font
    set_value_cell(report_ws, start_row + 4, column=4, value=normalize(get_value(row_data, "성별")), font=content_font)
    report_ws.cell(row=start_row + 5, column=1, value="기수").font = header_font
    set_value_cell(report_ws, start_row + 5, column=2, value=normalize(get_value(row_data, "기수")), font=content_font)
    report_ws.cell(row=start_row + 5, column=3, value="생년월일").font = header_font
    set_value_cell(report_ws, start_row + 5, column=4, value=normalize(get_value(row_data, "생년월일")), font=content_font)
    
    report_ws.cell(row=start_row + 6, column=1, value="교인구분").font = header_font
    set_value_cell(report_ws, start_row + 6, column=2, value=normalize(get_value(row_data, "교인구분")), font=content_font)
    
    report_ws.cell(row=start_row + 7, column=1, value="25 소속").font = header_font
    set_value_cell(report_ws, start_row + 7, column=2, value=normalize(get_value(row_data, "25 소속")), font=content_font)
    report_ws.cell(row=start_row + 7, column=3, value="25 직분").font = header_font
    set_value_cell(report_ws, start_row + 7, column=4, value=normalize(get_value(row_data, "25 직분")), font=content_font)
    
    current_status = get_value(row_data, "현재상태")
    current_status_note = get_value(row_data, "현재상태 기타설명란")
    planned = get_value(row_data, "예정사항")
    planned_note = get_value(row_data, "예정사항 기타설명란")

    current_status_value = current_status_note if current_status in (None, "") else current_status
    if current_status not in (None, "") and current_status_note not in (None, ""):
        current_status_value = f"{current_status}/{current_status_note}"

    planned_value = planned_note if planned in (None, "") else planned
    if planned not in (None, "") and planned_note not in (None, ""):
        planned_value = f"{planned}/{planned_note}"

    report_ws.cell(row=start_row + 8, column=1, value="현재상태").font = header_font
    set_value_cell(report_ws, start_row + 8, column=2, value=normalize(current_status_value), font=content_font)

    report_ws.cell(row=start_row + 9, column=1, value="예정사항").font = header_font
    set_value_cell(report_ws, start_row + 9, column=2, value=normalize(planned_value), font=content_font)
    
    report_ws.cell(row=start_row + 10, column=1, value="전화번호").font = header_font
    set_value_cell(report_ws, start_row + 10, column=2, value=normalize(get_value(row_data, "전화번호")), font=content_font)
    report_ws.cell(row=start_row + 11, column=1, value="토요예배 출석부 등급").font = header_font
    set_value_cell(report_ws, start_row + 11, column=2, value=normalize(get_value(row_data, "토요예배 출석부 등급")), font=content_font)
    
    # 출석 정보
    report_ws.cell(row=start_row + 12, column=1, value="출석정보").font = header_font
    report_ws.cell(row=start_row + 12, column=2, value="그룹모임").font = header_font
    set_value_cell(report_ws, start_row + 12, column=3, value=normalize(get_value(row_data, "그룹모임")), font=content_font)
    report_ws.cell(row=start_row + 12, column=4, value="주일낮").font = header_font
    set_value_cell(report_ws, start_row + 12, column=5, value=normalize(get_value(row_data, "주일낮예배")), font=content_font)
    report_ws.cell(row=start_row + 13, column=2, value="주일저녁").font = header_font
    set_value_cell(report_ws, start_row + 13, column=3, value=normalize(get_value(row_data, "주일저녁예배")), font=content_font)
    
    # 소견서
    opinion_text = get_text_or_fallback(row_data, "소견내용", "특별소견내용")
    report_ws.cell(row=start_row + 14, column=1, value="소견서").font = header_font
    report_ws.merge_cells(start_row=start_row + 14, start_column=2, end_row=start_row + 22, end_column=6)
    opinion_cell = report_ws.cell(row=start_row + 14, column=2, value=normalize(opinion_text))
    opinion_cell.font = opinion_font
    opinion_cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True, shrink_to_fit=True)
    for r in range(start_row + 14, start_row + 23):
        report_ws.row_dimensions[r].height = 28
    apply_row_border(report_ws, start_row + 14, 1, 1)
    for r in range(start_row + 14, start_row + 23):
        for c in range(2, 7):
            report_ws.cell(row=r, column=c).border = border
    
    # 비고
    report_ws.cell(row=start_row + 24, column=1, value="비고").font = header_font
    report_ws.merge_cells(start_row=start_row + 24, start_column=2, end_row=start_row + 26, end_column=6)
    set_value_cell(report_ws, start_row + 24, column=2, value=normalize(get_value(row_data, "비고")), font=content_font)
    report_ws.cell(row=start_row + 24, column=2).alignment = Alignment(horizontal='left', vertical='top')

    # 비고 다음 추가 정보 (삭제됨)
    
    # 소견서 담당자
    report_ws.cell(row=start_row + 29, column=1, value="담당자").font = header_font
    set_value_cell(report_ws, start_row + 29, column=2, value=normalize(get_value(row_data, "담당자")), font=content_font)
    report_ws.cell(row=start_row + 29, column=3, value="담당자 전화번호").font = header_font
    set_value_cell(report_ws, start_row + 29, column=4, value=normalize(get_value(row_data, "담당자 전화번호")), font=content_font)

    # 소견서 영역(14~22행)을 제외하고 행 단위 라인 적용
    rows_to_border = [4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 24, 25, 26, 29]
    for offset in rows_to_border:
        apply_row_border(report_ws, start_row + offset, 1, 7)

# 열 너비 자동 조정 (소견서 내용 제외)
ignore_width_cells = set()
for row_idx in range(1, len(data) + 1):
    start_row = (row_idx - 1) * 42 + 1
    ignore_width_cells.add((start_row + 14, 2))

for col in report_ws.columns:
    max_length = 0
    column = None
    for cell in col:
        if not isinstance(cell, openpyxl.cell.cell.MergedCell):
            if column is None:
                column = cell.column_letter
            try:
                if (cell.row, cell.column) in ignore_width_cells:
                    continue
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
    if column:
        report_ws.column_dimensions[column].width = (max_length + 2)

# 보고서 저장
output_file = 'report.xlsx'
report_wb.save(output_file)

print(f"보고서가 {output_file}로 저장되었습니다.")
